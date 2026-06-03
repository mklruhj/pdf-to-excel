"""
excel_writer.py — Build a styled .xlsx file from extracted rows.
"""
import io

import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

# ── Column definitions ─────────────────────────────────────────────────────────
COLUMNS = [
    "Item",
    "RFI",
    "Certificate_Code",
    "Issue_date",
    "Notes",
    "Quality Signature Name",
    "Quality Signature Date",
    "Client Signature Name",
    "Client Signature Date",
]

# ── Cell styles ────────────────────────────────────────────────────────────────
HEADER_FILL   = PatternFill("solid", fgColor="FFD966")   # yellow  (matches sample)
COMPLETE_FILL = PatternFill("solid", fgColor="C6EFCE")   # green
PARTIAL_FILL  = PatternFill("solid", fgColor="FFEB9C")   # amber
FAILED_FILL   = PatternFill("solid", fgColor="FFC7CE")   # red
NO_FILL       = PatternFill(fill_type=None)

THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)

HEADER_FONT = Font(name="Calibri", bold=True, size=11)
DATA_FONT   = Font(name="Calibri", size=11)
CENTER      = Alignment(horizontal="center", vertical="center", wrap_text=False)
LEFT        = Alignment(horizontal="left",   vertical="center", wrap_text=False)


def build_excel(rows: list) -> io.BytesIO:
    """
    Build a styled .xlsx file from a list of row dicts.
    Returns a BytesIO buffer ready for st.download_button.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "QC Register"
    ws.freeze_panes = "A2"   # freeze header row

    # ── Header row ──
    for ci, col_name in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=ci, value=col_name)
        cell.fill      = HEADER_FILL
        cell.font      = HEADER_FONT
        cell.alignment = CENTER
        cell.border    = THIN_BORDER

    # ── Data rows ──
    for ri, row in enumerate(rows, start=2):
        status = row.get("_status", "complete")
        fill   = (
            COMPLETE_FILL if status == "complete" else
            PARTIAL_FILL  if status == "partial"  else
            FAILED_FILL
        )

        for ci, col_name in enumerate(COLUMNS, start=1):
            value = row.get(col_name, "")
            cell  = ws.cell(row=ri, column=ci, value=value)
            cell.fill      = fill
            cell.font      = DATA_FONT
            cell.alignment = CENTER if col_name in ("Issue_date", "Quality Signature Date", "Client Signature Date") else LEFT
            cell.border    = THIN_BORDER

    # ── Column widths ──
    col_widths = {
        "Item":                   28,
        "RFI":                    32,
        "Certificate_Code":       28,
        "Issue_date":             14,
        "Notes":                  20,
        "Quality Signature Name": 24,
        "Quality Signature Date": 24,
        "Client Signature Name":  24,
        "Client Signature Date":  24,
    }
    for ci, col_name in enumerate(COLUMNS, start=1):
        ws.column_dimensions[
            openpyxl.utils.get_column_letter(ci)
        ].width = col_widths.get(col_name, 20)

    # ── Row height ──
    ws.row_dimensions[1].height = 22
    for ri in range(2, ws.max_row + 1):
        ws.row_dimensions[ri].height = 18

    # ── Save to buffer ──
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
